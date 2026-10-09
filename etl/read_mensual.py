"""Serie mensual del OPEX: primer semestre de cada año por origen y gran rubro, en USD y toneladas."""

from __future__ import annotations

import re
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook

from etl.aggregate import r1s
from etl.mappings import MEASURES, RUBROS, load_config, norm_name, put, vector

SEMESTER_LAST_MONTH = 6
REQUIRED = ("ano", "mes", "fob_dolar", "peso neto", "nombre prov", "rubro")


def read_mensual(path: Path) -> dict[str, dict[str, dict[str, list[float]]]]:
    cfg = load_config()
    origins = {**cfg["name_to_iso"], **cfg["series_origins"]}
    regions = cfg["regions"]
    series = {measure: defaultdict(lambda: defaultdict(vector)) for measure in MEASURES}
    workbook = load_workbook(path, read_only=True, data_only=True)
    rows = workbook.worksheets[0].iter_rows(values_only=True)
    index = _column_index(next(rows))
    for row in rows:
        if row[index["mes"]] is None or int(row[index["mes"]]) > SEMESTER_LAST_MONTH:
            continue
        origin = origins.get(norm_name(row[index["nombre prov"]]))
        if origin is None:
            raise ValueError(f"Origen desconocido en la serie mensual: {row[index['nombre prov']]!r}")
        rubro = _rubro(row[index["rubro"]])
        year = str(int(row[index["ano"]]))
        amounts = {"usd": float(row[index["fob_dolar"]] or 0) / 1e6, "tn": float(row[index["peso neto"]] or 0) / 1e6}
        keys = [origin, "N"]
        if origin in regions:
            keys.append("R:" + regions[origin])
        for measure, amount in amounts.items():
            for key in keys:
                put(series[measure][key][year], rubro, amount)
    workbook.close()
    return {
        measure: {key: {year: r1s(values) for year, values in sorted(years.items())} for key, years in by_key.items()}
        for measure, by_key in series.items()
    }


def _column_index(header: tuple) -> dict[str, int]:
    found = {norm_name(name): pos for pos, name in enumerate(header) if name is not None}
    missing = [name for name in REQUIRED if name not in found]
    if missing:
        raise ValueError(f"A la serie mensual le faltan columnas: {', '.join(missing)}")
    return found


def _rubro(label: object) -> int:
    match = re.search(r"\((\w+)\)", str(label))
    if not match or match.group(1) not in RUBROS:
        raise ValueError(f"Rubro desconocido en la serie mensual: {label!r}")
    return RUBROS.index(match.group(1)) + 1
